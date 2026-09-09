from __future__ import annotations

import os
import re
import shutil
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .backend_keywords import generate_backend_keywords
from . import listing_data
from .copy_writer import (
    BULLET_AUDIT_MAX_CHARS,
    DESCRIPTION_MAX_CHARS,
    ITEM_HIGHLIGHT_MAX_COUNT,
    TITLE_MAX_CHARS,
)
from .io import file_sha256, read_json, write_json
from .job import load_job
from .paths import FACTORY_ROOT
from .plugin import ProductPlugin
from .product_family import read_product_family
from .run_scope import read_run_scope, scoped_family_children
from .template_contract import template_input_fingerprint
from .template_field_plan import (
    compile_field_coverage,
    compile_field_plan,
    compile_field_requirements,
    country_of_origin_value,
    fit_template_image_slots as _fit_template_image_slots,
    field_expects_number,
    item_highlights_text,
    looks_numeric,
    primary_material_label,
    resolve_field,
    role_sort as _role_sort,
    split_weight_recommendation,
    template_allowed_values,
)
from .template_field_values import color_map_value, light_color_value, template_dimensions_from_facts
from .template_package import TemplatePackageError, write_plan_package
from .template_runtime import (
    DEFAULT_TEMPLATE_INVENTORY_AVAILABLE,
    DEFAULT_TEMPLATE_PRODUCT_ID_TYPE,
    child_sku,
    load_template_env,
    parent_sku,
    resolve_template_path,
    safe_component,
    template_job_with_defaults,
    template_mode as resolve_template_mode,
)
from .template_submit_ready import TemplateSubmitReadyError, assert_submit_ready_price_current
from .price import normalize_price_value
from .title_quality import listing_title_quality_issues
from .value_utils import truthy

DEFAULT_MARKETPLACE_ID = "ATVPDKIKX0DER"
MARKETPLACE_ID = os.environ.get("AMAZON_MARKETPLACE_ID", DEFAULT_MARKETPLACE_ID).strip() or DEFAULT_MARKETPLACE_ID
TEMPLATE_REFERENCE_MAX_ROWS = 12
TEMPLATE_REFERENCE_VALUE_LIMIT = 600
TEMPLATE_REFERENCE_EXAMPLE_LIMIT = 5
def _marketplace_field(value: str) -> str:
    return str(value).replace(f"marketplace_id={DEFAULT_MARKETPLACE_ID}", f"marketplace_id={MARKETPLACE_ID}")

IMAGE_FIELDS = {
    "main": "main_product_image_locator[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other1": "other_product_image_locator_1[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other2": "other_product_image_locator_2[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other3": "other_product_image_locator_3[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other4": "other_product_image_locator_4[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other5": "other_product_image_locator_5[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other6": "other_product_image_locator_6[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other7": "other_product_image_locator_7[marketplace_id=ATVPDKIKX0DER]#1.media_location",
    "other8": "other_product_image_locator_8[marketplace_id=ATVPDKIKX0DER]#1.media_location",
}

FIELD_ALIASES = {
    "sku": ["contribution_sku#1.value"],
    "child_sku": ["contribution_sku#1.value"],
    "product_type": ["product_type#1.value"],
    "record_action": ["::record_action"],
    "parentage": ["parentage_level[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "parent_sku": ["child_parent_sku_relationship[marketplace_id=ATVPDKIKX0DER]#1.parent_sku"],
    "variation_theme": ["variation_theme#1.name"],
    "item_name": ["item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "title": ["item_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "item_highlight": ["title_differentiation[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "brand": ["brand[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "brand_name": ["brand[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "product_id_type": ["amzn1.volt.ca.product_id_type"],
    "product_id": ["amzn1.volt.ca.product_id_value"],
    "item_type_keyword": ["item_type_keyword[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "model_number": ["model_number[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "model_name": ["model_name[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "manufacturer": ["manufacturer[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "description": ["product_description[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "bullet1": ["bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "bullet2": ["bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value"],
    "bullet3": ["bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value"],
    "bullet4": ["bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#4.value"],
    "bullet5": ["bullet_point[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#5.value"],
    "keywords": ["generic_keyword[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "material": ["material[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "material_type": ["material[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "fabric_type": ["fabric_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "color": ["color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "color_map": ["color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.standardized_values#1"],
    "color_name": ["color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "size": ["size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "size_name": ["size[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "style": ["style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "style_name": ["style[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "part_number": ["part_number[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "finish_type": ["finish_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "hardware_color": ["hardware_color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "light_color": ["light_color[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "special_feature": ["special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "special_feature1": ["special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "special_feature2": ["special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value"],
    "special_feature3": ["special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value"],
    "special_feature4": ["special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#4.value"],
    "special_feature5": ["special_feature[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#5.value"],
    "included_component": ["included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "included_component1": ["included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "included_component2": ["included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value"],
    "included_component3": ["included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value"],
    "included_component4": ["included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#4.value"],
    "included_component5": ["included_components[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#5.value"],
    "target_audience": ["target_audience[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "target_audience1": ["target_audience[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "target_audience2": ["target_audience[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value"],
    "target_audience3": ["target_audience[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value"],
    "target_audience4": ["target_audience[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#4.value"],
    "target_audience5": ["target_audience[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#5.value"],
    "mounting_type": ["mounting_type[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "plant_or_animal_product_type": ["plant_or_animal_product_type[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "number_of_items": ["number_of_items[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "package_quantity": ["item_package_quantity[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "unit_count": ["unit_count[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "unit_type": ["unit_count[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US].value"],
    "condition": ["condition_type[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "country": ["country_of_origin[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "dg": ["supplier_declared_dg_hz_regulation[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "batteries_required": ["batteries_required[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "batteries_included": ["batteries_included[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "contains_battery": ["contains_battery_or_cell[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "is_fragile": ["is_fragile[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "list_price": ["list_price[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "quantity": ["fulfillment_availability#1.quantity"],
    "fulfillment": ["fulfillment_availability#1.fulfillment_channel_code"],
    "inventory_available": ["fulfillment_availability#1.is_inventory_available"],
    "shipping_template": ["merchant_shipping_group[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "item_depth": ["item_depth[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "item_depth_unit": ["item_depth[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.unit"],
    "length": ["item_length_width_height[marketplace_id=ATVPDKIKX0DER]#1.length.value", "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.normalized_value.value", "item_display_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.value", "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.depth.value", "item_length[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "length_unit": ["item_length_width_height[marketplace_id=ATVPDKIKX0DER]#1.length.unit", "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.normalized_value.unit", "item_display_dimensions[marketplace_id=ATVPDKIKX0DER]#1.length.unit", "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.depth.unit", "item_length[marketplace_id=ATVPDKIKX0DER]#1.unit"],
    "width": ["item_length_width_height[marketplace_id=ATVPDKIKX0DER]#1.width.value", "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.normalized_value.value", "item_display_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.value", "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.width.value"],
    "width_unit": ["item_length_width_height[marketplace_id=ATVPDKIKX0DER]#1.width.unit", "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.normalized_value.unit", "item_display_dimensions[marketplace_id=ATVPDKIKX0DER]#1.width.unit", "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.width.unit"],
    "height": [
        "item_length_width_height[marketplace_id=ATVPDKIKX0DER]#1.height.value",
        "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.normalized_value.value",
        "item_display_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.value",
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.height.value",
    ],
    "height_unit": [
        "item_length_width_height[marketplace_id=ATVPDKIKX0DER]#1.height.unit",
        "item_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.normalized_value.unit",
        "item_display_dimensions[marketplace_id=ATVPDKIKX0DER]#1.height.unit",
        "item_depth_width_height[marketplace_id=ATVPDKIKX0DER]#1.height.unit",
    ],
    "item_weight": ["item_weight[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "item_weight_unit": ["item_weight[marketplace_id=ATVPDKIKX0DER]#1.unit"],
    "item_package_weight": ["item_package_weight[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "item_package_weight_unit": ["item_package_weight[marketplace_id=ATVPDKIKX0DER]#1.unit"],
    "load_capacity": ["weight_capacity[marketplace_id=ATVPDKIKX0DER]#1.maximum#1.value"],
    "load_capacity_unit": ["weight_capacity[marketplace_id=ATVPDKIKX0DER]#1.maximum#1.unit"],
    "max_weight": ["maximum_weight_recommendation[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "max_weight_unit": ["maximum_weight_recommendation[marketplace_id=ATVPDKIKX0DER]#1.unit"],
    "maximum_weight_recommendation": ["maximum_weight_recommendation[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "maximum_weight_recommendation_unit": ["maximum_weight_recommendation[marketplace_id=ATVPDKIKX0DER]#1.unit"],
    "number_of_doors": ["number_of_doors[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "number_of_drawers": ["number_of_drawers[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "number_of_shelves": ["number_of_shelves[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "shape": ["item_shape[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "frame_material": [
        "frame_material[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value",
        "frame[marketplace_id=ATVPDKIKX0DER]#1.material[language_tag=en_US]#1.value",
    ],
    "door_orientation": ["door[marketplace_id=ATVPDKIKX0DER]#1.orientation[language_tag=en_US]#1.value"],
    "room_type": ["room_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "recommended_uses": ["recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "recommended_use1": ["recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "recommended_use2": ["recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value"],
    "recommended_use3": ["recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value"],
    "recommended_use4": ["recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#4.value"],
    "recommended_use5": ["recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#5.value"],
    "room_type1": ["room_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "room_type2": ["room_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#2.value"],
    "room_type3": ["room_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#3.value"],
    "room_type4": ["room_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#4.value"],
    "room_type5": ["room_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#5.value"],
    "indoor_outdoor": ["recommended_uses_for_product[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "base_type": ["base_type[marketplace_id=ATVPDKIKX0DER][language_tag=en_US]#1.value"],
    "pot_material": ["container[marketplace_id=ATVPDKIKX0DER]#1.type[language_tag=en_US]#1.value"],
    "has_lights": ["has_builtin_light[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "has_led": ["has_builtin_light[marketplace_id=ATVPDKIKX0DER]#1.value"],
    "main_image_url": [IMAGE_FIELDS["main"]],
}

IMAGE_FIELDS = {key: _marketplace_field(value) for key, value in IMAGE_FIELDS.items()}
FIELD_ALIASES = {key: [_marketplace_field(value) for value in values] for key, values in FIELD_ALIASES.items()}

@dataclass(frozen=True)
class TemplateInfo:
    path: Path
    sheet: str
    max_row: int
    max_col: int
    fields: dict[str, int]
    labels: dict[str, str]
    required: list[dict[str, str]]
    existing_rows: list[dict[str, Any]]
    allowed_values: dict[str, list[str]] | None = None
    template_product_type: str = ""
    data_start_row: int = 7

class TemplateEngineError(RuntimeError):
    pass

def _assert_submit_ready_job_contract(job: dict[str, Any], *, mode: str) -> None:
    if mode != "submit_ready":
        return
    if not isinstance(job.get("gtin_exempt"), bool):
        raise TemplateEngineError("submit_ready job requires gtin_exempt to resolve to a boolean")
    if job["gtin_exempt"] is False and (
        not str(job.get("product_id_type") or "").strip()
        or not str(job.get("product_id") or "").strip()
    ):
        raise TemplateEngineError("non-exempt submit_ready job requires product_id_type and product_id")

def _template_scope_children(job_path: Path, family: dict[str, Any], *, mode: str) -> list[dict[str, Any]]:
    scope = read_run_scope(job_path)
    if mode == "submit_ready":
        family_count = len(family.get("family", {}).get("children") or [])
        selected = scope.get("selected_children") if isinstance(scope.get("selected_children"), list) else []
        if scope.get("mode") != "production" or len(selected) != family_count:
            raise TemplateEngineError("submit_ready template requires a full-family production RunScopeV5")
    children = scoped_family_children(family, job_path)
    if not children:
        raise TemplateEngineError("Template RunScopeV5 selected no ProductFamilyV3 children")
    return children

def build_template_plan(
    *,
    job_dir: str | Path,
    plugin: ProductPlugin,
    config_path: str = "",
    template_path: str = "",
    output: str = "",
    write_excel: bool = False,
    template_mode: str = "",
    artifact_root: str | Path = "",
) -> dict[str, Any]:
    job_path = Path(job_dir)
    job = load_job(job_path)
    env = load_template_env(
        str(FACTORY_ROOT / "config.env"),
        str(FACTORY_ROOT / "config.local.env"),
        config_path,
        job.get("config_path", ""),
    )
    try:
        mode = resolve_template_mode(env, template_mode)
    except ValueError as exc:
        raise TemplateEngineError(str(exc)) from exc
    family = read_product_family(job_path)
    if family.get("category_id") != plugin.category_id:
        raise TemplateEngineError("ProductFamilyV3 category does not match active template plugin")
    job = template_job_with_defaults(job, env=env)
    _assert_submit_ready_job_contract(job, mode=mode)
    if mode == "submit_ready":
        child_fetch_errors = list((family.get("source") or {}).get("child_fetch_errors") or [])
        if child_fetch_errors:
            raise TemplateEngineError(
                f"Submit-ready template requires the complete Apify family; {len(child_fetch_errors)} child fetch task(s) are unresolved"
            )
        from .publish import assert_published_release_complete

        assert_published_release_complete(job_dir=job_path, plugin=plugin)
    source_children = _template_scope_children(job_path, family, mode=mode)
    artifact_dir = Path(artifact_root) if str(artifact_root or "").strip() else job_path / "template"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    preflight = bool(str(artifact_root or "").strip())
    if not template_path and not output and not preflight:
        from .template_contract import template_plan_current

        if template_plan_current(
            job_dir=job_path,
            plugin=plugin,
            config_path=config_path,
            mode=mode,
            write_excel=write_excel,
        ):
            return read_json(job_path / "template" / "plan.json")
    try:
        template_file = resolve_template_path(
            plugin,
            job,
            factory_root=FACTORY_ROOT,
            explicit=template_path,
        )
    except (ValueError, FileNotFoundError) as exc:
        raise TemplateEngineError(str(exc)) from exc
    template = analyze_template(template_file)
    template_reference = _template_reference_report(template)
    published_images, image_audit = _published_image_plan(job_path, plugin=plugin)
    parent_asin = str(family["family"]["parent_asin"])
    output_path = Path(output) if output else artifact_dir / f"filled_{plugin.product_type}_{job['seed_asin']}.xlsm"
    theme = _theme_for_template(family["family"].get("variation_theme") or "", family["family"].get("variation_dimensions", []))
    rows: list[dict[str, Any]] = []
    listing_rows_by_sku: dict[str, dict[str, Any]] = {}
    children = source_children
    child_skus = [child_sku(job, child, index) for index, child in enumerate(children, start=1)]
    parent_sku_value = parent_sku(job, parent_asin, child_skus)
    copy_audit: list[dict[str, str]] = list(image_audit)
    copy_v1 = _load_copy_v1(
        job_path,
        env=env,
        mode=mode,
        children=children,
        plugin=plugin,
        job=job,
    )
    parent_specific = _parent_family_values(children, plugin)
    parent_child = {"asin": parent_asin, "variation_values": {}, "final_images": []}
    parent_listing = _build_listing_row_data(
        plugin=plugin,
        job=job,
        child=parent_child,
        specific=parent_specific,
        row_type="Parent",
        sku=parent_sku_value,
        asin=parent_asin,
        parent_sku="",
        theme=theme,
        env=env,
        copy_audit=copy_audit,
        copy_v1=copy_v1,
        template_mode=mode,
        published_images={},
    )
    parent_field_plan = compile_template_field_plan(
        template=template, plugin=plugin, row_data=parent_listing, template_mode=mode,
    )
    listing_rows_by_sku[parent_sku_value] = parent_listing
    rows.append({"row_type": "Parent", "sku": parent_sku_value, "source_asin": parent_asin, **parent_field_plan})

    for index, child in enumerate(children, start=1):
        asin = str(child.get("asin") or f"child-{index}")
        sku = child_sku(job, child, index)
        specific = _family_specific(child, plugin)
        child_listing = _build_listing_row_data(
            plugin=plugin,
            job=job,
            child=child,
            specific=specific,
            row_type="Child",
            sku=sku,
            asin=asin,
            parent_sku=parent_sku_value,
            theme=theme,
            env=env,
            copy_audit=copy_audit,
            copy_v1=copy_v1,
            template_mode=mode,
            published_images=published_images.get(asin, {}),
        )
        child_field_plan = compile_template_field_plan(
            template=template, plugin=plugin, row_data=child_listing, template_mode=mode,
        )
        listing_rows_by_sku[sku] = child_listing
        rows.append(
            {
                "row_type": "Child",
                "sku": sku,
                "source_asin": asin,
                "variation": child_listing.get("variation", {}),
                **child_field_plan,
            }
        )

    plan = {
        "template_engine": "template_field_plan",
        "template_path": str(template.path),
        "output_path": str(output_path),
        "reference_asin": job["seed_asin"],
        "job_id": job_path.name,
        "parent_asin": parent_asin,
        "category_id": plugin.category_id,
        "product_type": plugin.product_type,
        "variation_theme": theme,
        "template_input_fingerprint": template_input_fingerprint(
            job_path=job_path, template_path=template_file, output_path=output_path,
            mode=mode, write_excel=write_excel, job=job,
            plugin_config=plugin.merged_config(),
        ),
        "template": {
            "sheet": template.sheet,
            "max_col": template.max_col,
            "data_start_row": template.data_start_row,
            "field_count": len(template.fields),
            "required": template.required,
            "existing_rows": template.existing_rows,
        },
        "rows": rows,
        "audit": list(copy_audit),
    }
    listing_package = _build_listing_data_package(
        plan=plan,
        listing_rows_by_sku=listing_rows_by_sku,
        source_child_count=len(source_children),
        source_children=source_children,
        job=job,
        # Draft/preflight validates deterministic listing fields before paid
        # generation.  Published-image completeness remains a submit-ready
        # release gate, not a preflight gate.
        require_main_image=(
            IMAGE_FIELDS["main"] in template.fields
            and mode == "submit_ready"
        ),
    )
    _assert_submit_ready_price_current(listing_package.get("children", []), source_children, mode=mode)
    listing_data_path = artifact_dir / "listing_data.json"
    template_reference_path = artifact_dir / "template_reference.json"
    template_coverage_path = artifact_dir / "template_field_coverage.json"
    plan["listing_data_path"] = str(listing_data_path)
    plan["listing_data_summary"] = {
        "schema_version": listing_package["schema_version"],
        "source_child_count": listing_package["source_child_count"],
        "uploadable_child_count": listing_package["uploadable_child_count"],
        "omitted_child_count": len(listing_package.get("omitted_children", [])),
    }
    template_coverage = compile_field_coverage(
        fields=template.fields,
        labels=template.labels,
        rows=plan["rows"],
        reference=template_reference,
    )
    plan["template_reference_path"] = str(template_reference_path)
    plan["template_reference_summary"] = {
        "row_count": template_reference["row_count"],
        "field_example_count": template_reference["field_example_count"],
        "policy": template_reference["policy"],
    }
    plan["template_field_coverage_path"] = str(template_coverage_path)
    plan["template_field_coverage"] = template_coverage
    plan["template_mode"] = mode
    plan["copy_v1_path"] = str(copy_v1["_artifact_path"])
    plan["audit"].extend(listing_package.get("audit", []))
    plan["audit"].extend(_template_coverage_audit(template_coverage))
    plan["audit"].extend(_audit_plan(template, plan["rows"], job, plugin, mode=mode, source_children=source_children))
    blocking_audit = [
        item for item in plan["audit"]
        if isinstance(item, dict) and str(item.get("severity") or "").casefold() == "error"
    ]
    plan["artifact_state"] = (
        "submit_ready_blocked" if mode == "submit_ready" and blocking_audit
        else "submit_ready" if mode == "submit_ready"
        else "draft_with_blockers" if blocking_audit
        else "draft"
    )
    plan_path = artifact_dir / "plan.json"
    audit_path = artifact_dir / "audit.md"
    write_json(listing_data_path, listing_package)
    write_json(template_reference_path, template_reference)
    write_json(template_coverage_path, template_coverage)
    write_json(plan_path, plan)
    _write_markdown_audit(plan, audit_path)
    if mode == "submit_ready":
        _raise_on_template_audit_errors(plan)
    if write_excel:
        written = write_plan_workbook(plan, output_path)
        plan["workbook_sha256"] = file_sha256(written)
        if mode == "submit_ready":
            plan["final_template_path"] = str(_export_final_template(written, parent_asin, job_path.name))
            plan["final_template_sha256"] = file_sha256(Path(plan["final_template_path"]))
        else:
            plan["draft_template_path"] = str(written)
    write_json(plan_path, plan)
    _write_markdown_audit(plan, audit_path)
    return plan

def analyze_template(path: str | Path) -> TemplateInfo:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as exc:
        raise TemplateEngineError("openpyxl is required to analyze Amazon .xlsm templates") from exc

    template_path = Path(path)
    if not template_path.exists():
        raise TemplateEngineError(f"Template file not found: {template_path}")
    wb = load_workbook(template_path, read_only=False, keep_vba=False, data_only=False)
    try:
        if "Template" not in wb.sheetnames:
            raise TemplateEngineError(f"Workbook has no Template sheet: {template_path}")
        ws = wb["Template"]
        fields: dict[str, int] = {}
        labels: dict[str, str] = {}
        for col in range(1, ws.max_column + 1):
            field = ws.cell(5, col).value
            if field in (None, ""):
                continue
            text = str(field)
            fields[text] = col
            labels[text] = str(ws.cell(4, col).value or "")
        required = _read_required_definitions(wb)
        template_product_type = _read_template_product_type(wb)
        existing_rows = []
        sku_field = "contribution_sku#1.value"
        sku_col = fields.get(sku_field, 1)
        data_start_row = _detect_template_data_start_row(ws, sku_col)
        field_by_col = {col: field for field, col in fields.items()}
        for row in range(data_start_row, ws.max_row + 1):
            sku = ws.cell(row, sku_col).value
            values: dict[str, str] = {}
            for col, field in field_by_col.items():
                value = ws.cell(row, col).value
                if value not in (None, ""):
                    values[field] = _string_value(value)
            if values:
                existing_rows.append(
                    {
                        "row": row,
                        "sku": str(sku or values.get(sku_field) or ""),
                        "nonempty_field_count": len(values),
                        "values": values,
                    }
                )
        return TemplateInfo(
            template_path,
            "Template",
            ws.max_row,
            ws.max_column,
            fields,
            labels,
            required,
            existing_rows,
            template_allowed_values(wb, ws, fields, data_start_row, template_product_type),
            template_product_type,
            data_start_row,
        )
    finally:
        _close_workbook(wb)

def _detect_template_data_start_row(ws: Any, sku_col: int) -> int:
    settings = str(ws.cell(1, 1).value or "")
    declared = re.search(r"(?:^|&)dataRow=(\d+)(?:&|$)", settings)
    if declared:
        row_index = int(declared.group(1))
        while row_index <= int(ws.max_row or row_index):
            if not any(ws.cell(row_index, col).value not in (None, "") for col in range(1, int(ws.max_column or 1) + 1)):
                return row_index
            row_index += 1
        return row_index
    marker = str(ws.cell(7, sku_col).value or ws.cell(7, 1).value or "").strip().lower()
    if marker and (
        "prefilled attributes" in marker
        or "preference profiles" in marker
        or "please do not delete this row" in marker
    ):
        saw_prefill = False
        for row_index in range(8, min(int(ws.max_row or 8), 50) + 2):
            has_values = any(ws.cell(row_index, col).value not in (None, "") for col in range(1, int(ws.max_column or 1) + 1))
            if has_values:
                saw_prefill = True
                continue
            if saw_prefill or row_index == 8:
                return row_index
        return max(8, int(ws.max_row or 7) + 1)
    raise TemplateEngineError(
        "Cannot determine the Amazon template data start row from Amazon metadata or the protected prefilled-row marker; "
        "refusing to overwrite an unknown worksheet layout"
    )

def _read_template_product_type(wb: Any) -> str:
    try:
        defined_name = wb.defined_names.get("product_type1.value")
    except Exception:
        defined_name = None
    if defined_name is None:
        return ""
    try:
        destinations = list(defined_name.destinations)
    except Exception:
        return ""
    for sheet_name, coord in destinations:
        try:
            cells = wb[sheet_name][coord]
        except Exception:
            continue
        for cell in _iter_openpyxl_cells(cells):
            value = getattr(cell, "value", None)
            if value not in (None, ""):
                return str(value).strip()
    return ""

def _iter_openpyxl_cells(cells: Any):
    if hasattr(cells, "value"):
        yield cells
        return
    try:
        for item in cells:
            yield from _iter_openpyxl_cells(item)
    except TypeError:
        return

def write_plan_workbook(plan: dict[str, Any], output_path: str | Path) -> Path:
    try:
        return write_plan_package(plan, output_path, value_transform=_workbook_cell_value)
    except TemplatePackageError as exc:
        raise TemplateEngineError(str(exc)) from exc

def _close_workbook(workbook: Any) -> None:
    vba_archive = getattr(workbook, "vba_archive", None)
    if vba_archive is not None:
        try:
            vba_archive.close()
        except Exception:
            try:
                vba_archive.fp = None
            except Exception:
                pass
        try:
            workbook.vba_archive = None
        except Exception:
            pass
    close = getattr(workbook, "close", None)
    if callable(close):
        close()

def _export_final_template(path: Path, parent_asin: str, job_id: str = "") -> Path:
    target_dir = FACTORY_ROOT / "Final_templates" / safe_component(parent_asin)
    target_dir.mkdir(parents=True, exist_ok=True)
    suffix = f"_{safe_component(job_id)}" if job_id else ""
    target = target_dir / f"{path.stem}{suffix}{path.suffix}"
    tmp = target_dir / f".{target.name}.{os.getpid()}.tmp"
    try:
        shutil.copy2(path, tmp)
        os.replace(tmp, target)
    except PermissionError:
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
        target = target_dir / f"{path.stem}{suffix}_{timestamp}{path.suffix}"
        tmp = target_dir / f".{target.name}.{os.getpid()}.tmp"
        shutil.copy2(path, tmp)
        os.replace(tmp, target)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass
    return target

def _workbook_cell_value(field: str, value: Any) -> Any:
    if value in (None, ""):
        return value
    if field_expects_number(field) and looks_numeric(value):
        text = str(value).strip()
        if re.fullmatch(r"[+-]?\d+", text):
            try:
                return int(text)
            except ValueError:
                return value
        try:
            return float(text)
        except ValueError:
            return value
    return value

def resolve_template_field(template: TemplateInfo, alias: str) -> str | None:
    try:
        return resolve_field(template.fields, template.labels, FIELD_ALIASES, alias)
    except ValueError as exc:
        raise TemplateEngineError(str(exc)) from exc

def _build_listing_row_data(
    *,
    plugin: ProductPlugin,
    job: dict[str, Any],
    child: dict[str, Any],
    specific: dict[str, Any],
    row_type: str,
    sku: str,
    asin: str = "",
    parent_sku: str,
    theme: str,
    env: dict[str, str],
    copy_audit: list[dict[str, str]] | None = None,
    copy_v1: dict[str, Any],
    template_mode: str = "draft",
    published_images: dict[str, Any] | None = None,
) -> dict[str, Any]:
    raw_variation = child.get("variation_values", {}) if isinstance(child, dict) else {}
    variation = _normalized_variation_values_for_template(child, specific, raw_variation)
    title, item_highlights, bullets, description = _copy_v1_for_row(
        copy_v1,
        row_type=row_type,
        sku=sku,
        asin=asin,
    )
    if copy_audit is not None:
        copy_audit.append({"severity": "info", "sku": sku, "field": "copy", "message": "Using current CopyV1 listing copy"})
    material = _first(specific, "material", "frame_material", default="")
    country = str(job.get("country") or env.get("DEFAULT_COUNTRY_OF_ORIGIN") or env.get("AMAZON_FACTORY_COUNTRY_OF_ORIGIN") or "")
    list_price = _child_list_price(child, job, allow_job_override=template_mode != "submit_ready")
    color = str(variation.get("color") or variation.get("color_name") or specific.get("color") or "")
    size = str(variation.get("size") or variation.get("size_name") or _template_safe_variation_size(_first(specific, "size", "size_name", "size_class")) or "")
    style = str(variation.get("style") or variation.get("style_name") or specific.get("style") or "")
    model = str(specific.get("model_number") or "")
    item_count = specific.get("sold_unit_count") if row_type == "Child" else ""
    package_quantity = specific.get("package_quantity") if row_type == "Child" else ""
    attributes = dict(specific)
    attributes.update(
        {
            "material": material,
            "color": color,
            "size": size,
            "style": style,
            "model_number": model,
            "item_count": item_count,
            "package_quantity": package_quantity,
            "item_type_keyword": _item_type_keyword(plugin),
        }
    )
    parent_row = str(row_type).casefold() == "parent"
    return listing_data.row(
        row_type=row_type,
        sku=sku,
        asin=str(asin or (child.get("asin") if isinstance(child, dict) else "") or job.get("seed_asin") or sku),
        parent_sku=parent_sku if row_type == "Child" else "",
            variation=variation,
        copy={
            "title": title,
            "item_highlights": item_highlights,
            "bullets": bullets,
            "description": description,
        },
        images=(
            published_images
            if isinstance(published_images, dict) and published_images.get("main")
            else {"main": "", "other": []}
        ),
        attributes=attributes,
        # A parent is a relationship row, never a sellable offer. Keeping
        # offer fields on it creates invalid parent inventory/price records.
        offer={
            "list_price": "" if parent_row else list_price,
            "quantity": "" if parent_row else job.get("quantity", ""),
            "fulfillment": "" if parent_row else job.get("fulfillment", ""),
            "inventory_available": "" if parent_row else job.get("inventory_available", DEFAULT_TEMPLATE_INVENTORY_AVAILABLE),
            "shipping_template": "" if parent_row else job.get("shipping_template", ""),
        },
        compliance={
            "brand": job.get("brand", ""),
            "manufacturer": job.get("manufacturer", ""),
            "country": country,
            "condition": "" if parent_row else job.get("condition", ""),
            "dangerous_goods": job.get("dangerous_goods", ""),
            "batteries_required": "" if parent_row else job.get("batteries_required", ""),
            "batteries_included": "" if parent_row else job.get("batteries_included", ""),
            "contains_battery": "" if parent_row else job.get("contains_battery", ""),
            "is_fragile": "" if parent_row else job.get("is_fragile", ""),
            "gtin_exempt": bool(job.get("gtin_exempt", True)),
            "product_id_type": str(job.get("product_id_type") or DEFAULT_TEMPLATE_PRODUCT_ID_TYPE),
            "product_id": "" if job.get("gtin_exempt", True) or row_type != "Child" else str(job.get("product_id") or ""),
            "product_type": plugin.product_type,
            "variation_theme": theme,
            "parentage": row_type,
        },
    )

def _normalized_variation_values_for_template(
    child: dict[str, Any],
    specific: dict[str, Any],
    raw_variation: dict[str, Any] | None,
) -> dict[str, str]:
    normalized = child.get("normalized_facts") if isinstance(child, dict) and isinstance(child.get("normalized_facts"), dict) else {}
    normalized_variation = normalized.get("variation") if isinstance(normalized.get("variation"), dict) else {}
    variation = {str(key): str(value).strip() for key, value in {**(raw_variation or {}), **normalized_variation}.items() if str(value or "").strip()}
    specs = child.get("specs") if isinstance(child, dict) and isinstance(child.get("specs"), dict) else {}
    color = (
        str(normalized.get("color") or "")
        or _first(variation, "color", "color_name", "colour")
        or _first(specific, "color", "color_name", "colour")
        or _first(specs, "color", "color_name", "colour")
    )
    size_candidates = (
        _first(variation, "size", "size_name"),
        _first(specific, "size", "size_name", "size_class"),
        _first(specs, "size", "size_name", "size_class"),
        str(normalized.get("size") or ""),
    )
    size = next((_template_safe_variation_size(value) for value in size_candidates if _template_safe_variation_size(value)), "")
    style = str(normalized.get("style") or "") or _first(variation, "style", "style_name") or _first(specific, "style", "style_name") or _first(specs, "style", "style_name")
    if color:
        variation.setdefault("color", color)
        variation.setdefault("color_name", color)
    if size:
        variation.setdefault("size", size)
        variation.setdefault("size_name", size)
    if style:
        variation.setdefault("style", style)
        variation.setdefault("style_name", style)
    return variation

def _template_safe_variation_size(value: str) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.search(r"\d", text) and re.search(r"(?:\"|in|inch|cm|mm|ft|feet|[x×])", text, flags=re.I):
        return ""
    return text


def _plant_species_candidate(attributes: dict[str, Any]) -> str:
    evidence_values = [
        attributes.get("plant_or_animal_product_type"),
        attributes.get("tree_type"),
        attributes.get("plant_type"),
    ]
    for value in evidence_values:
        text = str(value or "").strip()
        if text:
            return text
    return ""


def _field_hints(value: Any) -> list[str]:
    text = str(value or "")
    return [part for part in re.split(r"[^A-Za-z0-9]+", text) if len(part) > 2]


def compile_template_field_plan(
    *, template: TemplateInfo, plugin: ProductPlugin, row_data: dict[str, Any],
    template_mode: str = "draft",
) -> dict[str, Any]:
    copy = row_data.get("copy") if isinstance(row_data.get("copy"), dict) else {}
    attributes = row_data.get("attributes") if isinstance(row_data.get("attributes"), dict) else {}
    compliance = row_data.get("compliance") if isinstance(row_data.get("compliance"), dict) else {}
    offer = row_data.get("offer") if isinstance(row_data.get("offer"), dict) else {}
    bullets = copy.get("bullets") if isinstance(copy.get("bullets"), list) else []
    title = str(copy.get("title") or "")
    description = str(copy.get("description") or "")
    sku = str(row_data.get("sku") or "")
    row_type = str(row_data.get("row_type") or "")
    parent_sku = str(row_data.get("parent_sku") or "")
    material_value = attributes.get("material", "")
    color_value = attributes.get("color", "")
    tree_type = _plant_species_candidate(attributes)
    item_type_candidate = attributes.get("item_type_keyword") or _item_type_keyword(plugin)
    semantic_hints = {
        "record_action": ["Create or Replace", "Update", "Create"],
        "brand": [str(compliance.get("brand") or "")],
        "item_type_keyword": _item_type_allowed_hints(plugin.category_id),
        "material": [str(material_value or ""), *_field_hints(material_value)],
        "color_map": [color_map_value(color_value, category_id=plugin.category_id), str(color_value or "")],
        "style": [str(attributes.get(key) or "") for key in ("source_title", "title", "style")],
        "country": ["China", "CN"],
        "inventory_available": ["Enabled", "Yes", "True", "Available", "Disabled", "No", "False", "Unavailable"],
        "recommended_uses": [str(attributes.get("recommended_uses") or attributes.get("recommended_use") or ""), "Indoor", "Outdoor", "Home", "Decor"],
        "plant_or_animal_product_type": [
            str(attributes.get(key) or "")
            for key in ("plant_or_animal_product_type", "tree_type", "plant_type", "source_title", "title", "model_name")
        ],
    }
    pairs = {
        "sku": sku,
        "product_type": template.template_product_type or compliance.get("product_type") or plugin.product_type,
        "record_action": "(Default) Create or Replace",
        "parentage": compliance.get("parentage") or row_type,
        "variation_theme": compliance.get("variation_theme") or "",
        "item_name": title,
        "brand": compliance.get("brand") or "",
        "product_id_type": compliance.get("product_id_type") or "",
        "product_id": compliance.get("product_id") or "",
        "item_type_keyword": item_type_candidate,
        "model_number": attributes.get("model_number") or "",
        "model_name": attributes.get("model_name") or "",
        "manufacturer": compliance.get("manufacturer") or "",
        "description": description,
        "keywords": generate_backend_keywords(
            category_id=plugin.category_id,
            title=title,
            bullets=bullets,
            description=description,
            product_specific=attributes,
            competitor_keywords=list(plugin.merged_config().get("competitor_keywords", [])),
            extra_terms=list(plugin.merged_config().get("backend_search_terms", [])),
            brand=str(compliance.get("brand") or ""),
        ),
        # Never coerce an unknown material into a convenient dropdown value.
        # A wrong material is worse than a missing one and must be audited.
        "material": material_value,
        "fabric_type": attributes.get("fabric_type") or primary_material_label(material_value),
        "color": color_value,
        "color_map": color_map_value(color_value, category_id=plugin.category_id),
        "size": attributes.get("size", ""),
        "style": attributes.get("style", ""),
        "part_number": attributes.get("part_number") or "",
        "finish_type": attributes.get("finish_type", ""),
        "hardware_color": attributes.get("hardware_color") or "",
        "light_color": attributes.get("light_color") or light_color_value(plugin.category_id, attributes),
        "mounting_type": attributes.get("mounting_type", ""),
        "plant_or_animal_product_type": tree_type,
        "number_of_items": attributes["item_count"],
        "package_quantity": attributes["package_quantity"],
        "unit_count": attributes["item_count"],
        "unit_type": "Count" if attributes.get("item_count") not in (None, "") else "",
        "condition": compliance.get("condition") or "",
        "country": country_of_origin_value(compliance.get("country")),
        "dg": compliance.get("dangerous_goods") or "",
        "batteries_required": compliance.get("batteries_required") or "",
        "batteries_included": compliance.get("batteries_included") or "",
        "contains_battery": compliance.get("contains_battery") or "",
        "is_fragile": compliance.get("is_fragile") or "",
        "list_price": offer.get("list_price", ""),
        "quantity": offer.get("quantity", ""),
        "fulfillment": offer.get("fulfillment", ""),
        "inventory_available": _inventory_available_candidate(offer.get("inventory_available", "")),
        "shipping_template": offer.get("shipping_template", ""),
        "recommended_uses": attributes.get("recommended_uses") or attributes.get("recommended_use") or "",
        "special_feature": attributes.get("special_feature") or "",
        "special_feature1": attributes.get("special_feature1") or attributes.get("special_feature") or "",
        "special_feature2": attributes.get("special_feature2") or "",
        "special_feature3": attributes.get("special_feature3") or "",
        "special_feature4": attributes.get("special_feature4") or "",
        "special_feature5": attributes.get("special_feature5") or "",
        "included_component": attributes.get("included_component") or attributes.get("included_components") or "",
        "target_audience": attributes.get("target_audience") or "",
        "number_of_doors": attributes.get("number_of_doors") or "",
        "number_of_drawers": attributes.get("number_of_drawers") or "",
        "number_of_shelves": attributes.get("number_of_shelves") or attributes.get("shelves") or "",
        "room_type": attributes.get("room_type") or "",
        "shape": attributes.get("shape") or "",
        "door_orientation": attributes.get("door_orientation") or "",
        "base_type": attributes.get("base_type") or "",
        "pot_material": attributes.get("pot_material") or "",
        "has_lights": attributes.get("has_lights") or attributes.get("has_led") or "",
        "parent_sku": parent_sku if row_type == "Child" else "",
        "frame_material": attributes.get("frame_material") or attributes.get("material", ""),
        "item_depth": attributes.get("item_depth") or attributes.get("depth") or "",
    }
    semantic_values = dict(pairs)
    semantic_values["item_highlight"] = item_highlights_text(copy)
    for index, bullet in enumerate(bullets[:5], start=1):
        semantic_values[f"bullet{index}"] = bullet
    weight_value, weight_unit = split_weight_recommendation(
        attributes.get("max_weight") or attributes.get("maximum_weight_recommendation"),
        attributes.get("max_weight_unit") or attributes.get("maximum_weight_recommendation_unit"),
    )
    semantic_values["max_weight"] = weight_value
    semantic_values["max_weight_unit"] = weight_unit
    for alias in ("length", "width", "height", "item_depth", "item_weight"):
        if attributes.get(alias):
            semantic_values[alias] = attributes.get(alias)
    for alias in ("length_unit", "width_unit", "height_unit", "item_depth_unit", "item_weight_unit"):
        base = alias.replace("_unit", "")
        if attributes.get(base):
            semantic_values[alias] = attributes.get(alias) or _default_unit_for(base)
    return compile_field_plan(
        fields=template.fields,
        labels=template.labels,
        aliases=FIELD_ALIASES,
        allowed_values=template.allowed_values or {},
        semantic_hints=semantic_hints,
        semantic_values=semantic_values,
        images=(row_data.get("images") if isinstance(row_data.get("images"), dict) else {}) if row_type == "Child" else {},
        image_fields=IMAGE_FIELDS,
        requirements=compile_field_requirements(
            fields=template.fields,
            labels=template.labels,
            aliases=FIELD_ALIASES,
            template_requirements=template.required,
            template_mode=template_mode,
        ),
        row_context={
            "row_type": row_type,
            "variation_theme": compliance.get("variation_theme") or "",
            "gtin_exempt": bool(compliance.get("gtin_exempt", True)),
            "contains_battery": compliance.get("contains_battery"),
            "batteries_required": compliance.get("batteries_required"),
            "dangerous_goods": compliance.get("dangerous_goods") or "",
            "package_level": attributes.get("package_level") or attributes.get("package_contains_sku") or "",
        },
    )

def _default_unit_for(base: str) -> str:
    # Weight units are not safely inferable from a bare number. Dimensions use
    # the US template's inch convention; a weight without an explicit unit is
    # left for audit instead of being mislabeled as Pounds.
    return "" if base == "item_weight" else "Inches"

def _inventory_available_candidate(value: Any) -> str:
    text = str(value or "").strip().casefold()
    if not text:
        return ""
    if text in {"1", "true", "yes", "y", "available", "enabled"}:
        return "Enabled"
    if text in {"0", "false", "no", "n", "unavailable", "disabled"}:
        return "Disabled"
    return str(value or "").strip()

def _listing_images_from_field_values(values: dict[str, Any]) -> tuple[str, list[str]]:
    main = str(values.get(IMAGE_FIELDS["main"]) or "")
    other: list[str] = []
    for index in range(1, 9):
        url = str(values.get(IMAGE_FIELDS[f"other{index}"]) or "")
        if url and url not in other:
            other.append(url)
    return main, other

def _build_listing_data_package(
    *,
    plan: dict[str, Any],
    listing_rows_by_sku: dict[str, dict[str, Any]],
    source_child_count: int,
    source_children: list[dict[str, Any]],
    job: dict[str, Any],
    require_main_image: bool,
) -> dict[str, Any]:
    kept_child_skus: set[str] = set()
    synced_rows: dict[str, dict[str, Any]] = {}
    for row in plan.get("rows", []):
        if not isinstance(row, dict):
            continue
        sku = str(row.get("sku") or "")
        row_data = listing_rows_by_sku.get(sku)
        if not row_data:
            continue
        values = row.get("field_values") if isinstance(row.get("field_values"), dict) else {}
        main, other = _listing_images_from_field_values(values)
        synced = listing_data.sync_images(row_data, main=main, other=other)
        synced_rows[sku] = synced
        if row.get("row_type") == "Child":
            kept_child_skus.add(sku)
    for sku, row_data in listing_rows_by_sku.items():
        synced_rows.setdefault(sku, row_data)
    parent_sku = next((str(row.get("sku") or "") for row in plan.get("rows", []) if isinstance(row, dict) and row.get("row_type") == "Parent"), "")
    if not parent_sku:
        parent_sku = next((sku for sku, row_data in synced_rows.items() if row_data.get("row_type") == "Parent"), "")
    source_child_rows: list[dict[str, Any]] = []
    for index, child in enumerate(source_children, start=1):
        sku = child_sku(job, child, index)
        row_data = synced_rows.get(sku) or listing_rows_by_sku.get(sku)
        if row_data:
            source_child_rows.append(row_data)
            continue
        variation = child.get("variation_values", {}) if isinstance(child, dict) else {}
        source_child_rows.append(
            listing_data.row(
                row_type="Child",
                sku=sku,
                asin=str(child.get("asin") or sku) if isinstance(child, dict) else sku,
                parent_sku=parent_sku,
                variation=variation if isinstance(variation, dict) else {},
            )
        )
    if not source_child_rows:
        source_child_rows = [row_data for row_data in synced_rows.values() if row_data.get("row_type") == "Child"]
    source_child_rows.sort(key=lambda item: str(item.get("asin") or item.get("sku") or ""))
    return listing_data.package(
        job_id=str(plan.get("job_id") or ""),
        parent_asin=str(plan.get("parent_asin") or ""),
        category_id=str(plan.get("category_id") or ""),
        product_type=str(plan.get("product_type") or ""),
        variation_theme=str(plan.get("variation_theme") or ""),
        parent=synced_rows.get(parent_sku, {}),
        source_children=source_child_rows[:source_child_count] if source_child_count else source_child_rows,
        uploadable_child_skus=kept_child_skus,
        require_main_image=require_main_image,
    )

def _published_image_plan(job_path: Path, *, plugin: ProductPlugin) -> tuple[dict[str, dict[str, Any]], list[dict[str, str]]]:
    from .publish import current_published_release_rows
    from .release_manifest import approved_release_rows, build_release_manifest

    audit: list[dict[str, str]] = []
    release = build_release_manifest(job_dir=job_path, plugin=plugin)
    current_rows = current_published_release_rows(job_path, approved_release_rows(release))
    if not current_rows:
        return {}, [{"severity": "warning", "sku": "*", "field": "images", "message": "No current published release images found"}]
    rows_by_child: dict[str, list[dict[str, str]]] = {}
    for row in current_rows:
        rows_by_child.setdefault(row["child"], []).append(row)
    result: dict[str, dict[str, Any]] = {}
    for child, rows in rows_by_child.items():
        rows.sort(key=lambda item: _role_sort(item.get("role", "")))
        main = next((row for row in rows if row.get("role") == "main"), None)
        if main is None:
            audit.append({
                "severity": "warning",
                "sku": child,
                "field": "main_image",
                "message": "Published rows have no main role; source images remain draft references and scene is not promoted to main",
            })
            continue
        kept, omitted = _fit_template_image_slots([row for row in rows if row is not main])
        if omitted:
            audit.append({
                "severity": "warning",
                "sku": child,
                "field": "images",
                "message": "Template image slots omitted roles by priority: " + ", ".join(row.get("role", "unknown") for row in omitted),
            })
        result[child] = {"main": main["url"], "other": [row["url"] for row in kept]}
    return result, audit

def _template_reference_report(template: TemplateInfo) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    examples: dict[str, dict[str, Any]] = {}
    for item in template.existing_rows[:TEMPLATE_REFERENCE_MAX_ROWS]:
        values = item.get("values") if isinstance(item.get("values"), dict) else {}
        row_values: dict[str, str] = {}
        for field, value in values.items():
            text = _short(str(value), TEMPLATE_REFERENCE_VALUE_LIMIT)
            row_values[field] = text
            bucket = examples.setdefault(
                field,
                {
                    "field": field,
                    "label": template.labels.get(field, ""),
                    "column": template.fields.get(field),
                    "policy": _template_reference_field_policy(field),
                    "examples": [],
                },
            )
            if text and text not in bucket["examples"] and len(bucket["examples"]) < TEMPLATE_REFERENCE_EXAMPLE_LIMIT:
                bucket["examples"].append(text)
        rows.append(
            {
                "row": item.get("row"),
                "sku": item.get("sku", ""),
                "nonempty_field_count": item.get("nonempty_field_count", len(row_values)),
                "values": row_values,
            }
        )
    return {
        "policy": "reference_only_not_auto_copied",
        "row_count": len(template.existing_rows),
        "captured_row_count": len(rows),
        "field_example_count": len(examples),
        "rows": rows,
        "fields": dict(sorted(examples.items(), key=lambda pair: (int(pair[1].get("column") or 10_000), pair[0]))),
    }

def _template_reference_field_policy(field: str) -> str:
    name = str(field or "").lower()
    if name in {"contribution_sku#1.value", "::record_action"} or "parent_sku" in name or "product_id" in name:
        return "identity_reference_only"
    if "image_locator" in name:
        return "media_reference_only"
    if "item_name" in name or "bullet_point" in name or "product_description" in name or "generic_keyword" in name:
        return "copy_style_reference_only"
    if "brand" in name or "manufacturer" in name or "model_" in name or "part_number" in name:
        return "brand_model_reference_only"
    return "attribute_reference_only"

def _template_coverage_audit(coverage: dict[str, Any]) -> list[dict[str, str]]:
    audit: list[dict[str, str]] = []
    unfilled_required = int(coverage.get("unfilled_required_field_count") or 0)
    if unfilled_required:
        audit.append(
            {
                "severity": "error",
                "sku": "*",
                "field": "template_required_fields",
                "message": f"{unfilled_required} required template field(s) are not filled by current mapping.",
            }
        )
    business_missing = int(coverage.get("business_required_missing_count") or 0)
    if business_missing:
        audit.append({
            "severity": "error", "sku": "*", "field": "factory_business_required",
            "message": f"{business_missing} project-required field value(s) are missing; see template_field_coverage.json.",
        })
    invalid_values = int(coverage.get("invalid_value_count") or 0)
    if invalid_values:
        audit.append({
            "severity": "error", "sku": "*", "field": "template_allowed_values",
            "message": f"{invalid_values} template value(s) are invalid; see template_field_coverage.json.",
        })
    conditional_risks = int(coverage.get("conditional_risk_count") or 0)
    if conditional_risks:
        audit.append({
            "severity": "warning", "sku": "*", "field": "conditional_allowed_values",
            "message": f"{conditional_risks} conditional or allowed-value selection(s) require review.",
        })
    unfilled_reference = int(coverage.get("unfilled_reference_field_count") or 0)
    if unfilled_reference:
        audit.append(
            {
                "severity": "info",
                "sku": "*",
                "field": "template_reference",
                "message": f"{unfilled_reference} field(s) have examples in the updated template but are not auto-filled; see template_field_coverage.json.",
            }
        )
    optional_invalid_values = int(coverage.get("optional_invalid_value_count") or 0)
    if optional_invalid_values:
        audit.append({
            "severity": "warning",
            "sku": "*",
            "field": "template_optional_allowed_values",
            "message": f"{optional_invalid_values} optional template value(s) were omitted because no legal dropdown value matched.",
        })
    return audit

def _assert_submit_ready_price_current(
    listing_rows: list[dict[str, Any]],
    source_children: list[dict[str, Any]],
    *,
    mode: str,
) -> None:
    try:
        assert_submit_ready_price_current(listing_rows, source_children, mode=mode)
    except TemplateSubmitReadyError as exc:
        raise TemplateEngineError(str(exc)) from exc

def _audit_plan(
    template: TemplateInfo,
    rows: list[dict[str, Any]],
    job: dict[str, Any],
    plugin: ProductPlugin,
    *,
    mode: str,
    source_children: list[dict[str, Any]] | None = None,
) -> list[dict[str, str]]:
    audit: list[dict[str, str]] = []
    for index, child in enumerate(source_children or [], start=1):
        if not isinstance(child, dict):
            continue
        flags = {str(flag) for flag in child.get("risk_flags") or []}
        if "variation_values_missing" in flags:
            audit.append({
                "severity": "error" if mode == "submit_ready" else "warning",
                "sku": child_sku(job, child, index),
                "field": "variation_values",
                "message": "Apify source data has no usable variation values for this child; verify the Amazon variation data before submit-ready output",
            })
    for row in rows:
        values = row.get("field_values", {})
        if row.get("row_type") == "Parent":
            for field, value in values.items():
                lowered = str(field or "").lower()
                if "parent_sku" in lowered and str(value or "").strip():
                    audit.append({"severity": "error", "sku": row["sku"], "field": "parent_sku", "message": "Parent row must not populate child parent_sku relationship field"})
                if any(token in lowered for token in ("list_price", "fulfillment_availability", "merchant_shipping_group")) and str(value or "").strip():
                    audit.append({"severity": "error", "sku": row["sku"], "field": field, "message": "Parent row must not contain sellable offer fields"})
    audit.append({"severity": "info", "sku": "*", "field": "template", "message": f"Built with {plugin.category_id} generic template mapping"})
    return audit

def _description_quality_issues(value: Any) -> list[str]:
    text = re.sub(r"\s+", " ", str(value or "").replace("<br>", " ").replace("<br/>", " ")).strip()
    if not text:
        return []
    return ["incomplete_trailing_phrase"] if re.search(r"\b(?:and|or|with|for)\s*$", text, flags=re.I) else []

def _raise_on_template_audit_errors(plan: dict[str, Any]) -> None:
    errors = [
        item
        for item in plan.get("audit", [])
        if isinstance(item, dict) and str(item.get("severity") or "").strip().lower() == "error"
    ]
    if not errors:
        return
    preview = "; ".join(
        f"{item.get('sku') or '*'} {item.get('field') or ''}: {item.get('message') or ''}".strip()
        for item in errors[:5]
    )
    raise TemplateEngineError(f"Template audit has {len(errors)} blocking error(s); see template/audit.md. {preview}")

def _write_markdown_audit(plan: dict[str, Any], path: Path) -> None:
    lines = [
        f"# Amazon {plan['product_type']} Template Audit",
        "",
        f"- Template: `{plan['template_path']}`",
        f"- Output: `{plan['output_path']}`",
        f"- Reference ASIN: `{plan['reference_asin']}`",
        f"- Parent ASIN: `{plan['parent_asin']}`",
        f"- Variation Theme: `{plan['variation_theme']}`",
        f"- Template Reference Rows: {plan.get('template_reference_summary', {}).get('row_count', 0)}",
        f"- Filled Template Fields: {plan.get('template_field_coverage', {}).get('filled_field_count', 0)} / {plan.get('template_field_coverage', {}).get('total_fields', 0)}",
        "",
        "## Rows",
        "",
        "| Row Type | SKU | Source ASIN | Variation | Filled Fields |",
        "|---|---|---|---|---:|",
    ]
    for row in plan["rows"]:
        variation = ", ".join(f"{k}={v}" for k, v in row.get("variation", {}).items())
        lines.append(f"| {row['row_type']} | `{row['sku']}` | `{row['source_asin']}` | {variation} | {len(row.get('field_values', {}))} |")
    coverage = plan.get("template_field_coverage") if isinstance(plan.get("template_field_coverage"), dict) else {}
    if coverage:
        lines.extend(
            [
                "",
                "## Template Field Coverage",
                "",
                f"- Required fields not filled anywhere: {coverage.get('unfilled_required_field_count', 0)}",
                f"- Project-required child values missing: {coverage.get('business_required_missing_count', 0)}",
                f"- Conditional / allowed-value risks: {coverage.get('conditional_risk_count', 0)}",
                f"- Fields with reference examples but no generated value: {coverage.get('unfilled_reference_field_count', 0)}",
                f"- Unmapped field count: {coverage.get('unmapped_field_count', 0)}",
            ]
        )
        samples = coverage.get("unfilled_reference_fields") if isinstance(coverage.get("unfilled_reference_fields"), list) else []
        if samples:
            lines.extend(["", "| Column | Field | Label | Policy | Example |", "|---:|---|---|---|---|"])
            for item in samples[:12]:
                examples = item.get("examples") if isinstance(item.get("examples"), list) else []
                example = _short(str(examples[0] if examples else ""), 100)
                lines.append(
                    f"| {item.get('column') or ''} | `{item.get('field') or ''}` | {item.get('label') or ''} | {item.get('policy') or ''} | {example} |"
                )
    lines.extend(["", "## Audit", "", "| Severity | SKU | Field | Message |", "|---|---|---|---|"])
    for item in plan.get("audit", []):
        lines.append(f"| {item['severity']} | {item['sku']} | `{item['field']}` | {item['message']} |")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")

def _read_required_definitions(wb: Any) -> list[dict[str, str]]:
    if "Data Definitions" not in wb.sheetnames:
        return []
    ws = wb["Data Definitions"]
    required = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if not row or not row[1]:
            continue
        required_text = str(row[5] or "").strip()
        if required_text.lower() not in {"required", "conditionally required"}:
            continue
        required.append(
            {
                "group": str(row[0] or ""),
                "field": str(row[1] or ""),
                "label": str(row[2] or ""),
                "required": required_text,
                "description": str(row[3] or ""),
            }
        )
    return required

def _load_copy_v1(
    job_path: Path,
    *,
    env: dict[str, str],
    mode: str,
    children: list[Any],
    plugin: ProductPlugin,
    job: dict[str, Any],
) -> dict[str, Any]:
    from .copy_polish import copy_request_fingerprint, read_copy_artifact

    data = read_copy_artifact(job_path, require_complete=True)
    expected_fingerprint = copy_request_fingerprint(
        env, mode=mode, children=children, plugin=plugin, job=job
    )
    if data.get("copy_request_fingerprint") != expected_fingerprint:
        raise TemplateEngineError("Template requires current CopyV1; rerun the copy stage")
    rows = data.get("rows") if isinstance(data.get("rows"), dict) else {}
    expected = {child_sku(job, child, index) for index, child in enumerate(children, start=1)} | {"__parent__"}
    if set(rows) != expected:
        raise TemplateEngineError("CopyV1 row inventory does not match ProductFamilyV3")
    invalid = [key for key, row in rows.items() if not isinstance(row, dict) or not _copy_row_has_ai_provenance(row)]
    if invalid:
        raise TemplateEngineError("CopyV1 has incomplete provider provenance: " + ", ".join(invalid[:10]))
    data = dict(data)
    data["_artifact_path"] = str(job_path / "reports" / "copy_v1.json")
    return data

def _copy_row_has_ai_provenance(row: dict[str, Any]) -> bool:
    provider = str(row.get("provider") or "").strip().lower()
    model = str(row.get("model") or "").strip()
    request_fingerprint = str(row.get("request_fingerprint") or "").strip()
    return provider not in {"", "rule"} and bool(model) and bool(request_fingerprint)

def _copy_v1_for_row(
    data: dict[str, Any],
    *,
    row_type: str,
    sku: str,
    asin: str,
) -> tuple[str, list[str], list[str], str]:
    rows = data.get("rows")
    if not isinstance(rows, dict):
        raise TemplateEngineError("CopyV1 rows are invalid")
    keys = ["__parent__", f"parent:{asin}", sku] if row_type == "Parent" else [asin, sku]
    row = next((rows.get(key) for key in keys if isinstance(rows.get(key), dict)), None)
    if row is None:
        raise TemplateEngineError(f"Template requires CopyV1 row for {row_type} sku={sku} asin={asin}")
    if not _copy_row_has_ai_provenance(row):
        raise TemplateEngineError(f"Template requires AI CopyV1 provenance for {sku}")
    title = str(row.get("title") or "").strip()
    item_highlights = [str(item).strip() for item in row.get("item_highlights", [])] if isinstance(row.get("item_highlights"), list) else []
    if row_type == "Parent":
        # Parent highlights are optional during CopyV1 validation, but the
        # template has one 2-5 phrase field. Keep the first authored phrases
        # instead of letting an overlong parent row block template creation.
        item_highlights = item_highlights[:ITEM_HIGHLIGHT_MAX_COUNT]
    bullets = [str(item).strip() for item in row.get("bullets", [])]
    description = str(row.get("description") or "").strip()
    title_issues = listing_title_quality_issues(title, category="")
    description_issues = _description_quality_issues(description)
    if (
        not title or len(title) > TITLE_MAX_CHARS
        or title_issues
        or not item_highlights_text({"item_highlights": item_highlights})
        or len(bullets) != 5 or any(not item or len(item) > BULLET_AUDIT_MAX_CHARS for item in bullets)
        or not description or len(description) > DESCRIPTION_MAX_CHARS or description_issues
    ):
        raise TemplateEngineError(f"Template found incomplete CopyV1 row for {sku}")
    return title, item_highlights, bullets, description

def _common_family_values(children: list[dict[str, Any]], plugin: ProductPlugin) -> dict[str, Any]:
    rows = [_family_specific(child, plugin) for child in children]
    if not rows:
        return {}
    common: dict[str, Any] = {}
    for key, value in rows[0].items():
        if value in (None, "", [], {}):
            continue
        if all(key in row and _template_values_equal(row[key], value) for row in rows[1:]):
            common[key] = value
    return common


def _template_values_equal(left: Any, right: Any) -> bool:
    if isinstance(left, str) or isinstance(right, str):
        normalize = lambda value: re.sub(r"\s+", " ", str(value or "")).strip().casefold()
        return normalize(left) == normalize(right)
    return left == right

def _parent_family_values(children: list[dict[str, Any]], plugin: ProductPlugin) -> dict[str, Any]:
    values = _common_family_values(children, plugin)
    if not values.get("fabric_type"):
        materials = [str(_family_specific(child, plugin).get("material") or "") for child in children]
        labels = [primary_material_label(value) for value in materials if value]
        if labels and all(label == labels[0] for label in labels):
            values["fabric_type"] = labels[0]
    return values

def _family_specific(child: dict[str, Any], plugin: ProductPlugin) -> dict[str, Any]:
    facts = dict(child.get("specs") or {}) if isinstance(child.get("specs"), dict) else {}
    product_specific = child.get("product_specific") if isinstance(child.get("product_specific"), dict) else {}
    category_facts = product_specific.get(plugin.category_id)
    if isinstance(category_facts, dict):
        facts.update(category_facts)
    normalized = child.get("normalized_facts") if isinstance(child.get("normalized_facts"), dict) else {}
    normalized_variation = normalized.get("variation") if isinstance(normalized.get("variation"), dict) else {}
    for key in ("color", "size", "style"):
        if not facts.get(key):
            facts[key] = normalized.get(key) or normalized_variation.get(key) or ""
    for key in (
        "dimensions", "product_dimensions", "physical_dimensions",
        "length", "width", "height", "item_depth",
        "length_unit", "width_unit", "height_unit", "item_depth_unit",
        "item_weight", "item_weight_unit", "weight",
    ):
        if facts.get(key) in (None, "", [], {}) and normalized.get(key) not in (None, "", [], {}):
            facts[key] = normalized[key]
    facts["sold_unit_count"] = normalized.get("sold_unit_count") or child.get("sold_unit_count")
    facts["package_quantity"] = normalized.get("package_quantity") or child.get("package_quantity")
    resolved = normalized.get("resolved_facts") if isinstance(normalized.get("resolved_facts"), dict) else {}
    for fact_name, record in resolved.items():
        if not isinstance(record, dict):
            continue
        status = str(record.get("status") or "")
        value = record.get("value")
        if status == "confirmed" and value not in (None, "", []):
            facts[str(fact_name)] = value
        elif status == "conflicted":
            # A conflicting raw spec must never fall through to a template
            # column.  Keep the conflict in ProductFamily for audit/review.
            normalized_name = re.sub(r"[^a-z0-9]+", "_", str(fact_name).casefold()).strip("_")
            for raw_key in list(facts):
                if re.sub(r"[^a-z0-9]+", "_", str(raw_key).casefold()).strip("_") == normalized_name:
                    facts.pop(raw_key, None)
    facts = _normalized_template_fact_keys(facts)
    facts.update(template_dimensions_from_facts(facts, category_id=plugin.category_id))
    if not facts.get("frame_material") and facts.get("material"):
        facts["frame_material"] = facts["material"]
    return {str(key): value for key, value in facts.items() if value not in (None, "", [], {})}

def _normalized_template_fact_keys(facts: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(facts)
    for key, value in list(facts.items()):
        alias = re.sub(r"[^a-z0-9]+", "_", str(key).casefold()).strip("_")
        if alias and normalized.get(alias) in (None, "", [], {}):
            normalized[alias] = value
    return normalized

def _item_type_keyword(plugin: ProductPlugin) -> str:
    keywords = plugin.merged_config().get("competitor_keywords")
    if isinstance(keywords, list) and keywords:
        return str(keywords[0])
    return plugin.display_name.lower().replace("_", " ")


def _item_type_allowed_hints(category_id: str) -> list[str]:
    return {
        "artificial_tree": ["Artificial Trees (artificial-trees)"],
        "bed_frame": ["> Bed Frames (bed-frames)"],
        "bathroom_cabinet": ["Storage Cabinets (cabinets)"],
        "medicine_cabinet": ["Medical Cabinets (medicine-cabinets)"],
    }.get(str(category_id or ""), [])

def _child_list_price(child: dict[str, Any], job: dict[str, Any], *, allow_job_override: bool = True) -> str:
    # Price authority is the Apify source evidence (normalized facts, then the
    # child offer). job.list_price is a draft-only convenience and can never
    # fill submit-ready output.
    normalized = child.get("normalized_facts") if isinstance(child.get("normalized_facts"), dict) else {}
    normalized_price = normalize_price_value(normalized.get("list_price"))
    if normalized_price:
        return normalized_price
    offer = child.get("offer") if isinstance(child.get("offer"), dict) else {}
    offer_price = normalize_price_value(offer.get("list_price"))
    if offer_price:
        return offer_price
    if not allow_job_override:
        return ""
    return normalize_price_value(job.get("list_price"))

def _theme_for_template(theme: str, dimensions: list[str]) -> str:
    if theme:
        return theme
    parts: list[str] = []
    for dim in dimensions:
        normalized = _norm(dim)
        if "color" in normalized:
            parts.append("COLOR")
        elif "size" in normalized or "height" in normalized:
            parts.append("SIZE")
        elif "style" in normalized:
            parts.append("STYLE_NAME")
    # Do not invent a variation theme when ProductFamily has none. The
    # template's ontology must be supplied by Fetch, not guessed from fields.
    return "/".join(dict.fromkeys(parts)) if parts else ""

def _first(data: dict[str, Any], *keys: str, default: str = "") -> str:
    for key in keys:
        value = data.get(key)
        if value not in (None, "", []):
            return str(value)
    return default

def _string_value(value: Any) -> str:
    if isinstance(value, list):
        return ", ".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, dict):
        return ", ".join(f"{key}: {val}" for key, val in value.items() if val not in (None, ""))
    return str(value)

def _short(value: str, limit: int) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) <= limit:
        return text
    cut = text[: limit - 1].rstrip()
    best = -1
    for separator in (". ", "; ", ", ", " "):
        best = max(best, cut.rfind(separator))
    if best >= max(40, int(limit * 0.6)):
        cut = cut[:best].rstrip()
    cut = cut.rstrip(" ,;:-")
    cut = re.sub(r"\b(?:of|and|or|with|for|to|the|a|an)$", "", cut, flags=re.IGNORECASE).rstrip(" ,;:-")
    return f"{cut}." if cut else text[: limit - 1].rstrip() + "."

def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())

